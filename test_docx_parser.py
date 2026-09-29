import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.docx_parser import DocxSurveyParser
from app.models import QuestionType


class DocxParserTests(unittest.TestCase):
    def _build_sample(self, path: Path) -> None:
        document = Document()
        document.add_heading("Тестовая анкета", level=1)
        document.add_paragraph("S1. Какие продукты вы покупали?")
        document.add_paragraph("Несколько ответов")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "1"
        table.rows[0].cells[1].text = "Бургеры"
        row = table.add_row()
        row.cells[0].text = "99"
        row.cells[1].text = "Ничего из перечисленного"
        document.add_paragraph("Код 99 — исключающий")
        document.add_paragraph("Q1. Оцените идею")
        document.add_paragraph("Один ответ")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "1"
        table.rows[0].cells[1].text = "Не нравится"
        row = table.add_row()
        row.cells[0].text = "5"
        row.cells[1].text = "Очень нравится"
        document.add_paragraph("Q2. Почему идея вам нравится?")
        document.add_paragraph("Открытый ответ")
        document.add_paragraph("Показать Q2, если в Q1 выбран код 5")
        document.save(path)

    def test_reads_questions_options_types_and_logic(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.docx"
            self._build_sample(path)
            spec = DocxSurveyParser().parse(path)

        self.assertEqual([question.id for question in spec.questions], ["S1", "Q1", "Q2"])
        self.assertEqual(spec.questions[0].type, QuestionType.MULTI)
        self.assertTrue(spec.questions[0].options[-1].exclusive)
        self.assertEqual(spec.questions[2].type, QuestionType.TEXT)
        self.assertEqual(len(spec.questions[2].display_rules), 1)
        self.assertEqual(spec.questions[2].display_rules[0].condition.atoms[0].values, ["5"])

    def test_reads_uncoded_paragraph_options_and_ignores_formula(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "plain-options.docx"
            document = Document()
            document.add_paragraph("S1. Выберите ресторан")
            document.add_paragraph("Один ответ")
            document.add_paragraph("Бургер Кинг")
            document.add_paragraph("Другой ресторан")
            document.add_paragraph("E5. При какой цене блюдо покажется дорогим?")
            document.add_paragraph("E4 < E5 < E2")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        self.assertEqual([option.code for option in spec.questions[0].options], ["1", "2"])
        self.assertEqual(spec.questions[0].type, QuestionType.SINGLE)
        self.assertEqual(spec.questions[1].type, QuestionType.NUMBER)
        self.assertEqual(spec.questions[1].options, [])

    def test_infers_previous_question_for_code_only_screening_rule(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "screening.docx"
            document = Document()
            document.add_paragraph("S1. Выберите ресторан")
            document.add_paragraph("Один ответ")
            document.add_paragraph("Бургер Кинг")
            document.add_paragraph("Другой ресторан")
            document.add_paragraph("Закончить интервью, если не выбран код 1")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        rule = next(rule for rule in spec.rules if rule.action == "terminate")
        self.assertEqual(rule.condition.atoms[0].question_id, "S1")
        self.assertEqual(rule.condition.atoms[0].operator, "not_in")
        self.assertEqual(rule.condition.atoms[0].values, ["1"])

    def test_ignores_full_and_partial_strikethrough(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "strikethrough.docx"
            document = Document()
            question = document.add_paragraph()
            question.add_run("Q1. Какой ")
            removed = question.add_run("бургер ")
            removed.font.strike = True
            question.add_run("ролл Вам нравится?")
            document.add_paragraph("Один ответ")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "1"
            table.rows[0].cells[1].text = "Чикен ролл"
            row = table.add_row()
            row.cells[0].text = "2"
            removed_option = row.cells[1].paragraphs[0].add_run("Воппер")
            removed_option.font.strike = True

            removed_rule = document.add_paragraph()
            rule_run = removed_rule.add_run("Показать Q3, если в Q1 выбран код 1")
            rule_run.font.strike = True

            removed_question = document.add_paragraph()
            removed_question_run = removed_question.add_run("Q2. Удалённый вопрос?")
            removed_question_run.font.strike = True
            document.add_paragraph("Открытый ответ")
            document.add_paragraph("Этот текст относился к удалённому вопросу")
            document.add_paragraph("Q3. Активный вопрос?")
            document.add_paragraph("Открытый ответ")
            document.save(path)

            spec = DocxSurveyParser().parse(path)

        self.assertEqual([item.id for item in spec.questions], ["Q1", "Q3"])
        self.assertEqual(spec.questions[0].text, "Какой ролл Вам нравится?")
        self.assertEqual([item.text for item in spec.questions[0].options], ["Чикен ролл"])
        self.assertEqual(spec.rules, [])
        self.assertNotIn("Этот текст относился", " ".join(spec.questions[0].instructions))


if __name__ == "__main__":
    unittest.main()
